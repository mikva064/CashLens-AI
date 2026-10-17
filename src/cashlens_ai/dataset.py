"""Audit local video labels and split by video content, without copying footage.

The workbook's first three columns are filename, weak timestamp, binary label.
Numeric timestamps (including numeric strings) are interpreted as seconds, never
as Excel day fractions. A timestamp is an annotation point, not an event interval.
Only ``audit_dataset`` needs openpyxl; model inference does not import it.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import os
from pathlib import Path
import random
import re
import tempfile
from typing import Any
import unicodedata

from .errors import DatasetValidationError


_MANUAL_FIELDS = ("type", "start_sec", "end_sec", "rule_id", "evidence_sec", "review_notes")
_REVIEW_FIELDS = (
    "source_row", "label_filename", "filename", "sha256", "label",
    "timestamp_raw", "timestamp_sec", "annotation_kind", "split", "review_flags",
    *_MANUAL_FIELDS,
)
_UNITS = re.compile(
    r"(?P<number>\d+(?:[.,]\d+)?)\s*"
    r"(?P<unit>час(?:а|ов)?|ч|hours?|h|минут(?:а|ы)?|мин|м|minutes?|min|m|"
    r"секунд(?:а|ы)?|сек|с|seconds?|sec|s)\.?",
    re.IGNORECASE,
)


def parse_timestamp_seconds(value: Any) -> float | None:
    """Parse seconds, MM:SS/HH:MM:SS, or Russian/English time units.

    Empty cells mean no point annotation. A plain number always means seconds;
    e.g. ``0.5`` means half a second. Invalid values fail instead of disappearing.
    """
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        raise ValueError("boolean is not a timestamp")
    if isinstance(value, (int, float)):
        result = float(value)
    elif isinstance(value, str):
        text = value.strip().casefold().replace("\u00a0", " ")
        if re.fullmatch(r"\d+(?:[.,]\d+)?", text):
            result = float(text.replace(",", "."))
        elif re.fullmatch(r"\d+:\d{1,2}(?::\d{1,2})?(?:[.,]\d+)?", text):
            parts = [float(part.replace(",", ".")) for part in text.split(":")]
            if any(part >= 60 for part in parts[1:]):
                raise ValueError("clock timestamp has minutes or seconds >= 60")
            result = 0.0
            for part in parts:
                result = result * 60 + part
        else:
            matches = list(_UNITS.finditer(text))
            remaining = _UNITS.sub("", text)
            if not matches or remaining.strip(" ,;+\t\r\n"):
                raise ValueError(f"unsupported timestamp {value!r}")
            result = 0.0
            used_units: set[int] = set()
            for match in matches:
                unit = match["unit"]
                factor = 3600 if unit.startswith(("ч", "h")) else (
                    60 if unit.startswith(("м", "m")) else 1
                )
                if factor in used_units:
                    raise ValueError("a timestamp repeats a time unit")
                used_units.add(factor)
                result += float(match["number"].replace(",", ".")) * factor
    else:
        raise ValueError(f"unsupported timestamp type {type(value).__name__}")
    if not math.isfinite(result) or result < 0:
        raise ValueError("timestamp must be a finite nonnegative number of seconds")
    return result


def _parse_label(value: Any) -> int:
    if isinstance(value, (int, float)) and value in (0, 1):
        return int(value)
    if isinstance(value, str):
        normalized = " ".join(value.casefold().split())
        labels = {"0": 0, "1": 1, "нарушение": 1, "без нарушения": 0}
        if normalized in labels:
            return labels[normalized]
    raise ValueError(f"unknown binary label {value!r}; expected нарушение/без нарушения or 1/0")


def _name_key(name: str, *, confusable: bool = False) -> str:
    key = unicodedata.normalize("NFC", name).casefold()
    # Only the known X/Х filename prefix is a fallback. Never transliterate names.
    if confusable and key.startswith("х-"):
        key = "x-" + key[2:]
    return key


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _workbook_rows(labels_path: Path) -> tuple[str, list[dict[str, Any]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise DatasetValidationError(
            "Reading label workbooks requires openpyxl; install the project's data extra."
        ) from exc
    try:
        workbook = load_workbook(labels_path, read_only=True, data_only=True)
    except Exception as exc:
        raise DatasetValidationError(f"Cannot read labels workbook {labels_path}: {exc}") from exc
    try:
        sheet = workbook["labels"] if "labels" in workbook.sheetnames else workbook.active
        header = next(sheet.iter_rows(min_row=1, max_row=1, max_col=3, values_only=True))
        first_header = " ".join(str(header[0] or "").casefold().split())
        if first_header not in {"наименование файла", "имя файла", "файл", "filename", "file"}:
            raise DatasetValidationError(
                "Column A must have a filename header; expected наименование файла or filename."
            )
        records: list[dict[str, Any]] = []
        errors = []
        for row_num, values in enumerate(
            sheet.iter_rows(min_row=2, max_col=3, values_only=True), start=2
        ):
            filename, timestamp, label = values
            if all(value is None or (isinstance(value, str) and not value.strip()) for value in values):
                continue
            try:
                if not isinstance(filename, str) or not filename.strip():
                    raise ValueError("filename is empty or not text")
                filename = filename.strip()
                if "/" in filename or "\\" in filename or filename in {".", ".."}:
                    raise ValueError("filename must be a basename, not a path")
                requested = filename if filename.casefold().endswith(".mov") else filename + ".mov"
                parsed_label = _parse_label(label)
                timestamp_sec = parse_timestamp_seconds(timestamp)
                records.append({
                    "source_row": row_num,
                    "label_filename": filename,
                    "requested_filename": requested,
                    "label": parsed_label,
                    "timestamp_raw": str(timestamp) if timestamp is not None else "",
                    "timestamp_sec": timestamp_sec,
                    "annotation_kind": "weak_point" if timestamp_sec is not None else "video_label_only",
                })
            except ValueError as exc:
                errors.append(f"row {row_num}: {exc}")
        if errors:
            raise DatasetValidationError("Invalid labels:\n" + "\n".join(errors))
        if not records:
            raise DatasetValidationError("The labels workbook has no annotation rows.")
        return sheet.title, records
    finally:
        workbook.close()


def _atomic_text(path: Path, content: str, *, encoding: str = "utf-8") -> None:
    temporary: str | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="w", encoding=encoding, newline="", dir=path.parent, delete=False
        ) as stream:
            temporary = stream.name
            stream.write(content)
        os.replace(temporary, path)
    finally:
        if temporary is not None and os.path.exists(temporary):
            os.unlink(temporary)


def _review_csv(path: Path, records: list[dict[str, Any]]) -> str:
    # Re-running the audit must not erase a teammate's manual review work.
    existing: dict[tuple[str, str, str], dict[str, str]] = {}
    if path.exists():
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            for row in csv.DictReader(stream):
                key = (row.get("source_row", ""), row.get("filename", ""), row.get("sha256", ""))
                existing[key] = row
    buffer = io.StringIO(newline="")
    writer = csv.DictWriter(buffer, fieldnames=_REVIEW_FIELDS)
    writer.writeheader()
    for record in records:
        key = (str(record["source_row"]), record["filename"], record["sha256"])
        previous = existing.get(key, {})
        row = {field: record.get(field, "") for field in _REVIEW_FIELDS}
        row["review_flags"] = ";".join(record["review_flags"])
        for field in _MANUAL_FIELDS:
            row[field] = previous.get(field, "")
        writer.writerow(row)
    return buffer.getvalue()


def audit_dataset(
    labels_path: Path,
    videos_dir: Path,
    output_dir: Path,
    *,
    seed: int = 42,
    holdout_per_class: int = 5,
) -> dict[str, Any]:
    """Write manifest.json, review.csv and a stratified SHA256 group split.

    Filenames first match case-insensitively with their original alphabet. The
    Latin X/Cyrillic Х prefix is considered only when the exact match is absent;
    such matches are recorded for review. All top-level .mov files are hashed,
    including unused files. Duplicate content can belong to only one partition.
    Conflicting binary labels, absent videos and invalid cells prevent all writes.
    Manual CSV review columns survive subsequent audits of the same source rows.
    """
    labels_path, videos_dir, output_dir = map(Path, (labels_path, videos_dir, output_dir))
    if not labels_path.is_file():
        raise DatasetValidationError(f"Labels workbook does not exist: {labels_path}")
    if not videos_dir.is_dir():
        raise DatasetValidationError(f"Videos directory does not exist: {videos_dir}")
    if not isinstance(seed, int) or isinstance(seed, bool):
        raise DatasetValidationError("seed must be an integer")
    if not isinstance(holdout_per_class, int) or isinstance(holdout_per_class, bool) or holdout_per_class < 0:
        raise DatasetValidationError("holdout_per_class must be a nonnegative integer")
    sheet_name, records = _workbook_rows(labels_path)
    video_paths = sorted(
        (path for path in videos_dir.iterdir() if path.is_file() and path.suffix.casefold() == ".mov"),
        key=lambda path: (path.name.casefold(), path.name),
    )
    exact: dict[str, list[Path]] = {}
    fallback: dict[str, list[Path]] = {}
    for path in video_paths:
        exact.setdefault(_name_key(path.name), []).append(path)
        fallback.setdefault(_name_key(path.name, confusable=True), []).append(path)
    missing: list[str] = []
    ambiguous: list[str] = []
    for record in records:
        requested = record["requested_filename"]
        candidates = exact.get(_name_key(requested), [])
        match_kind = "exact"
        if not candidates:
            candidates = fallback.get(_name_key(requested, confusable=True), [])
            match_kind = "confusable_prefix"
        if not candidates:
            missing.append(f"row {record['source_row']}: {requested}")
        elif len(candidates) != 1:
            ambiguous.append(f"row {record['source_row']}: {requested}")
        else:
            record["filename"] = candidates[0].name
            record["filename_match"] = match_kind
    if missing or ambiguous:
        details = []
        if missing:
            details.append("Missing source videos:\n" + "\n".join(missing))
        if ambiguous:
            details.append("Ambiguous filenames:\n" + "\n".join(ambiguous))
        raise DatasetValidationError("\n".join(details))
    videos = [{"filename": path.name, "sha256": _sha256(path), "size_bytes": path.stat().st_size}
              for path in video_paths]
    by_filename = {video["filename"]: video for video in videos}
    groups: dict[str, dict[str, Any]] = {}
    for video in videos:
        group = groups.setdefault(video["sha256"], {
            "sha256": video["sha256"], "filenames": [], "source_rows": [],
            "label": None, "weak_points_sec": [], "review_flags": [], "split": "unused",
        })
        group["filenames"].append(video["filename"])
    for record in records:
        record["sha256"] = by_filename[record["filename"]]["sha256"]
        group = groups[record["sha256"]]
        if group["label"] is not None and group["label"] != record["label"]:
            raise DatasetValidationError(
                "Conflicting binary labels for duplicate content "
                f"{record['sha256']}: files {', '.join(group['filenames'])}; "
                "resolve the annotations before making a split."
            )
        group["label"] = record["label"]
        group["source_rows"].append(record["source_row"])
        if record["timestamp_sec"] is not None:
            group["weak_points_sec"].append(record["timestamp_sec"])
    rng = random.Random(seed)
    development: list[str] = []
    holdout: list[str] = []
    class_group_counts: dict[str, int] = {}
    for label in (0, 1):
        hashes = sorted(sha for sha, group in groups.items() if group["label"] == label)
        class_group_counts[str(label)] = len(hashes)
        if len(hashes) < holdout_per_class:
            raise DatasetValidationError(
                f"Class {label} has {len(hashes)} unique video groups; "
                f"cannot reserve {holdout_per_class} for holdout."
            )
        rng.shuffle(hashes)
        holdout.extend(hashes[:holdout_per_class])
        development.extend(hashes[holdout_per_class:])
    for partition, hashes in (("development", development), ("holdout", holdout)):
        for sha in hashes:
            groups[sha]["split"] = partition
    for group in groups.values():
        # Keep each source annotation above; unique points here are just a summary.
        group["weak_points_sec"] = sorted(set(group["weak_points_sec"]))
        if len(group["filenames"]) > 1:
            group["review_flags"].append("duplicate_content")
            if len(group["weak_points_sec"]) > 1:
                group["review_flags"].append("duplicate_timestamps_differ")
    for record in records:
        group = groups[record["sha256"]]
        record["split"] = group["split"]
        record["review_flags"] = list(group["review_flags"])
        if record["filename_match"] != "exact":
            record["review_flags"].append("confusable_prefix_match")
        if record["label"] == 1 and record["timestamp_sec"] is None:
            record["review_flags"].append("positive_without_timestamp")
        if record["label"] == 0 and record["timestamp_sec"] is not None:
            record["review_flags"].append("negative_has_timestamp")
    used_filenames = {record["filename"] for record in records}
    unused = [video["filename"] for video in videos if video["filename"] not in used_filenames]
    duplicate_groups = [group for group in groups.values() if len(group["filenames"]) > 1]
    summary: dict[str, Any] = {
        "annotation_rows": len(records),
        "positive_rows": sum(record["label"] == 1 for record in records),
        "negative_rows": sum(record["label"] == 0 for record in records),
        "video_files": len(videos),
        "labeled_video_files": len(used_filenames),
        "unique_labeled_videos": sum(group["label"] is not None for group in groups.values()),
        "unique_video_groups": len(groups),
        "unique_groups_per_class": class_group_counts,
        "duplicate_group_count": len(duplicate_groups),
        "duplicate_groups": duplicate_groups,
        "unused_video_files": unused,
        "missing_source_files": [],
        "development_groups": len(development),
        "holdout_groups": len(holdout),
        "review_required_rows": sum(bool(record["review_flags"]) for record in records),
    }
    manifest = {
        "schema_version": 1,
        "source": {"labels_path": str(labels_path.resolve()), "sheet": sheet_name,
                   "videos_dir": str(videos_dir.resolve())},
        "timestamp_policy": "Weak annotation points only; numeric values are seconds; no intervals inferred.",
        "summary": summary,
        "annotations": records,
        "videos": videos,
        "groups": [groups[sha] for sha in sorted(groups)],
    }
    split = {
        "schema_version": 1,
        "group_key": "sha256",
        "seed": seed,
        "holdout_per_class": holdout_per_class,
        "development": sorted(development),
        "holdout": sorted(holdout),
        "unused": sorted(sha for sha, group in groups.items() if group["label"] is None),
    }
    # Validate and assemble all content before creating any output files.
    review_text = _review_csv(output_dir / "review.csv", records)
    output_dir.mkdir(parents=True, exist_ok=True)
    _atomic_text(output_dir / "manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
    _atomic_text(output_dir / "split.json", json.dumps(split, ensure_ascii=False, indent=2) + "\n")
    _atomic_text(output_dir / "review.csv", review_text, encoding="utf-8-sig")
    return {**summary, "manifest_path": str((output_dir / "manifest.json").resolve()),
            "split_path": str((output_dir / "split.json").resolve()),
            "review_path": str((output_dir / "review.csv").resolve())}
